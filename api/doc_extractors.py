import os
import io
import csv
from pathlib import Path
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook


def extract_pdf_text(file_content: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(file_content))
        if not reader.pages:
            raise ValueError("PDF contains no pages.")
        text_parts = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                text_parts.append(text)
        return "\n\n".join(text_parts).strip()
    except Exception as e:
        raise ValueError(f"Unable to extract text from PDF: {str(e)}")


def extract_docx_text(file_content: bytes) -> str:
    try:
        document = Document(io.BytesIO(file_content))

        text_parts = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        # Also extract table content
        for table in document.tables:
            for row in table.rows:
                row_text = []

                for cell in row.cells:
                    row_text.append(cell.text.strip())

                text_parts.append(" | ".join(row_text))

        return "\n".join(text_parts).strip()

    except Exception as e:
        raise ValueError(f"Unable to extract text from DOCX: {str(e)}")


def extract_text_file(file_content: bytes) -> str:
    try:
        return file_content.decode("utf-8").strip()

    except UnicodeDecodeError:
        try:
            return file_content.decode("utf-8-sig").strip()

        except UnicodeDecodeError as e:
            raise ValueError("Text file is not valid UTF-8.") from e


def extract_csv_text(file_content: bytes) -> str:
    try:
        text = file_content.decode("utf-8-sig")

        reader = csv.reader(io.StringIO(text))

        rows = []

        for row in reader:
            rows.append(" | ".join(cell.strip() for cell in row))

        return "\n".join(rows).strip()

    except UnicodeDecodeError as e:
        raise ValueError("CSV file is not valid UTF-8.") from e

    except csv.Error as e:
        raise ValueError("CSV file is malformed.") from e


def extract_xlsx_text(file_content: bytes) -> str:
    try:
        workbook = load_workbook(
            filename=io.BytesIO(file_content),
            read_only=True,
            data_only=True,
        )

        text_parts = []

        for worksheet in workbook.worksheets:

            text_parts.append(f"Sheet: {worksheet.title}")

            for row in worksheet.iter_rows(values_only=True):
                values = []

                for value in row:
                    if value is not None:
                        values.append(str(value).strip())

                if values:
                    text_parts.append(" | ".join(values))

        workbook.close()

        return "\n".join(text_parts).strip()

    except Exception as e:
        raise ValueError(f"Unable to extract text from XLSX: {str(e)}")


def detect_file_type(
    file_content: bytes,
    filename: str,
) -> str:

    extension = Path(filename).suffix.lower()

    # PDF signature
    if file_content.startswith(b"%PDF"):
        return ".pdf"

    # DOCX/XLSX are ZIP-based formats
    if file_content.startswith(b"PK"):
        if extension == ".docx":
            return ".docx"

        if extension == ".xlsx":
            return ".xlsx"

        return extension

    # Plain text / CSV / Markdown
    if extension in {
        ".txt",
        ".md",
        ".csv",
    }:
        try:
            file_content.decode("utf-8-sig")
            return extension
        except UnicodeDecodeError:
            raise ValueError("File content does not match a valid text encoding.")

    raise ValueError("Unsupported or unrecognized file format.")


def extract_text(
    file_content: bytes,
    file_type: str,
) -> str:

    if file_type == ".pdf":
        return extract_pdf_text(file_content)

    elif file_type == ".docx":
        return extract_docx_text(file_content)

    elif file_type in {".txt", ".md"}:
        return extract_text_file(file_content)

    elif file_type == ".csv":
        return extract_csv_text(file_content)

    elif file_type == ".xlsx":
        return extract_xlsx_text(file_content)

    else:
        raise ValueError(f"No extractor available for {file_type}")
